from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import config_workbook
from config_workbook import (
    ConfigWorkbookError,
    ensure_editable_config,
    load_config_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "成本配置模板.xlsx"


class ConfigWorkbookTests(unittest.TestCase):
    def _copy_template(self, directory: Path, name: str = "config.xlsx") -> Path:
        target = directory / name
        shutil.copyfile(TEMPLATE, target)
        return target

    def test_template_contains_all_migrated_data(self) -> None:
        config = load_config_workbook(TEMPLATE, check_lock=False)

        self.assertEqual(4, len(config.price_data))
        self.assertEqual(17, sum(len(item["keywords"]) for item in config.price_data))
        self.assertEqual(
            89,
            sum(
                len(product["价格"])
                for item in config.price_data
                for product in item["products"]
            ),
        )
        self.assertEqual(
            48, sum(len(prices) for prices in config.pillow_cost_data.values())
        )
        self.assertEqual(9, len(config.moving_costs_data))
        self.assertEqual(2, len(config.others_cost_data))
        self.assertEqual(5.5, config.dropship_unit_cost)

    def test_new_pillow_row_is_loaded(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._copy_template(Path(temp_dir))
            workbook = load_workbook(path)
            worksheet = workbook["枕芯价格"]
            worksheet.append(["200*70", "测试棉", 12.5, 13, 13, 15, 16, 16])
            workbook.save(path)
            workbook.close()

            config = load_config_workbook(path, check_lock=False)
            self.assertEqual(12.5, config.pillow_cost_data["200*70"]["测试棉枕芯"])
            self.assertEqual(16, config.pillow_cost_data["200*70"]["加重kk测试棉枕芯"])

    def test_new_base_size_column_is_loaded(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._copy_template(Path(temp_dir))
            workbook = load_workbook(path)
            worksheet = workbook["基础商品价格"]
            new_column = worksheet.max_column + 1
            worksheet.cell(1, new_column, "70*200")
            for row in range(2, worksheet.max_row + 1):
                if (
                    worksheet.cell(row, 1).value == "等身枕套"
                    and worksheet.cell(row, 2).value == "桃皮绒"
                ):
                    worksheet.cell(row, new_column, 88)
                    break
            workbook.save(path)
            workbook.close()

            config = load_config_workbook(path, check_lock=False)
            category = next(
                item for item in config.price_data if item["category"] == "等身枕套"
            )
            product = next(
                item for item in category["products"] if item["尺寸"] == "70*200"
            )
            self.assertEqual(88, product["价格"]["桃皮绒"])

    def test_missing_sheet_reports_workbook_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._copy_template(Path(temp_dir))
            workbook = load_workbook(path)
            del workbook["其他成本"]
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ConfigWorkbookError, "缺少工作表"):
                load_config_workbook(path, check_lock=False)

    def test_partial_row_reports_sheet_row_and_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._copy_template(Path(temp_dir))
            workbook = load_workbook(path)
            worksheet = workbook["其他成本"]
            row_number = worksheet.max_row + 1
            worksheet.cell(row_number, 1, "测试")
            workbook.save(path)
            workbook.close()

            with self.assertRaises(ConfigWorkbookError) as context:
                load_config_workbook(path, check_lock=False)
            message = str(context.exception)
            self.assertIn("工作表「其他成本」", message)
            self.assertIn(f"第 {row_number} 行", message)
            self.assertIn("「单价」", message)

    def test_formula_duplicate_and_negative_price_are_rejected(self) -> None:
        mutations = (
            ("公式", lambda ws: ws.cell(2, 2, "=1+1"), "不允许使用公式"),
            (
                "重复",
                lambda ws: ws.append([ws.cell(2, 1).value, ws.cell(2, 2).value]),
                "重复",
            ),
            ("负数", lambda ws: ws.cell(2, 2, -1), "非负数字"),
        )
        for label, mutate, expected in mutations:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp_dir:
                path = self._copy_template(Path(temp_dir))
                workbook = load_workbook(path)
                worksheet = workbook["其他成本"]
                mutate(worksheet)
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(ConfigWorkbookError, expected):
                    load_config_workbook(path, check_lock=False)

    def test_invalid_version_header_size_merge_and_category_reference(self) -> None:
        def wrong_version(workbook):
            workbook["全局参数"].cell(2, 2, 2)

        def wrong_header(workbook):
            workbook["其他成本"].cell(1, 2, "价格")

        def wrong_size(workbook):
            workbook["枕芯价格"].cell(2, 1, "120x40")

        def merged_data(workbook):
            workbook["其他成本"].merge_cells("A2:B2")

        def unknown_category(workbook):
            workbook["基础商品关键词"].append(["不存在类别", "测试"])

        cases = (
            (wrong_version, "只支持配置版本 1"),
            (wrong_header, "表头必须严格为"),
            (wrong_size, "尺寸必须使用"),
            (merged_data, "数据区域不允许合并单元格"),
            (unknown_category, "未在「基础商品类别」中定义"),
        )
        for mutate, expected in cases:
            with self.subTest(expected=expected), tempfile.TemporaryDirectory() as temp_dir:
                path = self._copy_template(Path(temp_dir))
                workbook = load_workbook(path)
                mutate(workbook)
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(ConfigWorkbookError, expected):
                    load_config_workbook(path, check_lock=False)

    def test_excel_lock_file_prevents_loading(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            path = self._copy_template(directory, "成本配置.xlsx")
            (directory / "~$成本配置.xlsx").write_bytes(b"lock")
            with self.assertRaisesRegex(ConfigWorkbookError, "正在 Excel 中打开"):
                load_config_workbook(path)

    def test_ensure_editable_config_copies_once_without_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            with patch.object(config_workbook, "_application_dir", return_value=directory), patch.object(
                config_workbook, "_template_path", return_value=TEMPLATE
            ):
                destination = ensure_editable_config()
                self.assertTrue(destination.is_file())
                destination.write_bytes(b"user-data")
                self.assertEqual(destination, ensure_editable_config())
                self.assertEqual(b"user-data", destination.read_bytes())


if __name__ == "__main__":
    unittest.main()
