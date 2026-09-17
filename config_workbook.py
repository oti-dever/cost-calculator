"""成本配置工作簿的生命周期、加载与校验。

运行时始终读取可编辑的「成本配置.xlsx」。如果该文件不存在，
程序会从内嵌的「成本配置模板.xlsx」原子复制一份。
"""

from __future__ import annotations

import os
import re
import shutil
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.exceptions import InvalidFileException


CONFIG_WORKBOOK_FILENAME = "成本配置.xlsx"
CONFIG_TEMPLATE_FILENAME = "成本配置模板.xlsx"
CONFIG_VERSION = 1

INSTRUCTION_SHEET = "使用说明"
GLOBAL_SHEET = "全局参数"
BASE_CATEGORY_SHEET = "基础商品类别"
BASE_KEYWORD_SHEET = "基础商品关键词"
BASE_PRICE_SHEET = "基础商品价格"
PILLOW_PRICE_SHEET = "枕芯价格"
MOVING_PRICE_SHEET = "义乳义臀价格"
OTHER_PRICE_SHEET = "其他成本"

REQUIRED_SHEETS = (
    INSTRUCTION_SHEET,
    GLOBAL_SHEET,
    BASE_CATEGORY_SHEET,
    BASE_KEYWORD_SHEET,
    BASE_PRICE_SHEET,
    PILLOW_PRICE_SHEET,
    MOVING_PRICE_SHEET,
    OTHER_PRICE_SHEET,
)

EXPECTED_HEADERS: Dict[str, Tuple[str, ...]] = {
    GLOBAL_SHEET: ("参数", "值", "说明"),
    BASE_CATEGORY_SHEET: ("类别名称", "优先级"),
    BASE_KEYWORD_SHEET: ("类别名称", "关键词"),
    MOVING_PRICE_SHEET: ("关键词", "单价"),
    OTHER_PRICE_SHEET: ("关键词", "单价"),
}

BASE_PRICE_FIXED_HEADERS = ("类别名称", "材质")
PILLOW_MATRIX_HEADERS = (
    "尺寸",
    "枕芯材质",
    "标准等身",
    "标准分腿",
    "标准开孔",
    "加重等身",
    "加重分腿",
    "加重开孔",
)
PILLOW_STYLE_PREFIXES = {
    "标准等身": "",
    "标准分腿": "ft",
    "标准开孔": "kk",
    "加重等身": "加重",
    "加重分腿": "加重ft",
    "加重开孔": "加重kk",
}

SIZE_PATTERN = re.compile(r"^\d+\*\d+$")
DROPSHIP_KEYWORDS = ("枕芯", "yr", "义乳", "yt", "义臀")


class ConfigWorkbookError(ValueError):
    """配置工作簿不可用或数据不合法。"""


@dataclass
class CostConfig:
    """经过完整校验、可供计算引擎直接使用的配置。"""

    price_data: list
    moving_costs_data: list
    pillow_cost_data: dict
    others_cost_data: list
    dropship_unit_cost: float
    dropship_keywords: Tuple[str, ...]
    source_path: Path


def _application_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def _template_path() -> Path:
    if getattr(sys, "frozen", False):
        bundle_dir = Path(getattr(sys, "_MEIPASS", _application_dir()))
        return bundle_dir / CONFIG_TEMPLATE_FILENAME
    return Path(__file__).resolve().parent / CONFIG_TEMPLATE_FILENAME


def get_editable_config_path() -> Path:
    """返回使用者可编辑的配置文件路径。"""

    return _application_dir() / CONFIG_WORKBOOK_FILENAME


def ensure_editable_config() -> Path:
    """
    确保运行目录中存在可编辑配置。

    仅在文件不存在时从模板复制，绝不覆盖已有配置。
    """

    destination = get_editable_config_path()
    if destination.exists():
        return destination

    template = _template_path()
    if not template.is_file():
        raise ConfigWorkbookError(f"找不到内嵌配置模板：{template}")

    temp_path = destination.with_name(
        f".{destination.name}.{uuid.uuid4().hex}.tmp"
    )
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(template, temp_path)
        os.replace(temp_path, destination)
    except OSError as exc:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise ConfigWorkbookError(
            f"无法在程序目录创建「{CONFIG_WORKBOOK_FILENAME}」，"
            f"请检查目录写入权限：{destination.parent}（{exc}）"
        ) from exc

    return destination


def _ensure_not_in_use(path: Path) -> None:
    """尽可能提前发现配置正被 Excel 占用的情况。"""

    excel_lock_file = path.with_name(f"~${path.name}")
    if excel_lock_file.exists():
        raise ConfigWorkbookError(
            f"配置文件正在 Excel 中打开：{path}。"
            "请保存并关闭配置文件后重试。"
        )

    try:
        with path.open("r+b"):
            pass
    except PermissionError as exc:
        raise ConfigWorkbookError(
            f"无法读写配置文件：{path}。"
            "请保存并关闭 Excel，或检查文件权限后重试。"
        ) from exc
    except OSError as exc:
        raise ConfigWorkbookError(f"无法访问配置文件：{path}（{exc}）") from exc


def _location(
    path: Path,
    sheet_name: str | None = None,
    row: int | None = None,
    column: str | None = None,
) -> str:
    parts = [f"{path.name} 配置错误"]
    if sheet_name:
        parts.append(f"工作表「{sheet_name}」")
    if row is not None:
        parts.append(f"第 {row} 行")
    if column:
        parts.append(f"「{column}」")
    return "，".join(parts)


def _raise_config_error(
    path: Path,
    message: str,
    sheet_name: str | None = None,
    row: int | None = None,
    column: str | None = None,
) -> None:
    raise ConfigWorkbookError(
        f"{_location(path, sheet_name, row, column)}：{message}"
    )


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean_text(
    value: Any,
    path: Path,
    sheet_name: str,
    row: int,
    column: str,
) -> str:
    if not isinstance(value, str):
        _raise_config_error(
            path,
            "必须填写文本",
            sheet_name,
            row,
            column,
        )
    text = value.strip()
    if not text:
        _raise_config_error(
            path,
            "不能为空",
            sheet_name,
            row,
            column,
        )
    return text


def _number(
    value: Any,
    path: Path,
    sheet_name: str,
    row: int,
    column: str,
    *,
    non_negative: bool = True,
) -> int | float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        _raise_config_error(
            path,
            "必须是数字",
            sheet_name,
            row,
            column,
        )
    number = value
    if non_negative and number < 0:
        _raise_config_error(
            path,
            "必须是非负数字",
            sheet_name,
            row,
            column,
        )
    return number


def _positive_integer(
    value: Any,
    path: Path,
    sheet_name: str,
    row: int,
    column: str,
) -> int:
    number = _number(value, path, sheet_name, row, column, non_negative=False)
    if int(number) != number or number <= 0:
        _raise_config_error(
            path,
            "必须是正整数",
            sheet_name,
            row,
            column,
        )
    return int(number)


def _size(
    value: Any,
    path: Path,
    sheet_name: str,
    row: int,
    column: str,
) -> str:
    text = _clean_text(value, path, sheet_name, row, column)
    if not SIZE_PATTERN.fullmatch(text):
        _raise_config_error(
            path,
            "尺寸必须使用「数字*数字」格式，例如 50*160",
            sheet_name,
            row,
            column,
        )
    return text


def _assert_no_data_merges(path: Path, worksheet) -> None:
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.max_row >= 2:
            _raise_config_error(
                path,
                f"数据区域不允许合并单元格（{merged_range}）",
                worksheet.title,
            )


def _assert_headers(path: Path, worksheet, expected: Sequence[str]) -> None:
    actual = [worksheet.cell(row=1, column=index).value for index in range(1, len(expected) + 1)]
    actual_clean = [value.strip() if isinstance(value, str) else value for value in actual]
    if tuple(actual_clean) != tuple(expected):
        _raise_config_error(
            path,
            f"表头必须严格为：{' | '.join(expected)}",
            worksheet.title,
            1,
        )

    for column in range(len(expected) + 1, worksheet.max_column + 1):
        if not _is_blank(worksheet.cell(row=1, column=column).value):
            _raise_config_error(
                path,
                "表头后存在未支持的额外列",
                worksheet.title,
                1,
            )
        for row in range(2, worksheet.max_row + 1):
            if not _is_blank(worksheet.cell(row=row, column=column).value):
                _raise_config_error(
                    path,
                    "未支持的额外列存在数据，但第 1 行没有表头",
                    worksheet.title,
                    row,
                    f"第{column}列",
                )


def _read_rows(path: Path, worksheet, headers: Sequence[str]) -> List[Tuple[int, Dict[str, Any]]]:
    _assert_no_data_merges(path, worksheet)
    _assert_headers(path, worksheet, headers)

    rows: List[Tuple[int, Dict[str, Any]]] = []
    for row_index in range(2, worksheet.max_row + 1):
        cells: List[Cell] = [
            worksheet.cell(row=row_index, column=column)
            for column in range(1, len(headers) + 1)
        ]
        values = [cell.value for cell in cells]
        if all(_is_blank(value) for value in values):
            continue

        for header, cell, value in zip(headers, cells, values):
            if cell.data_type == "f" or (
                isinstance(value, str) and value.startswith("=")
            ):
                _raise_config_error(
                    path,
                    "数据单元格不允许使用公式",
                    worksheet.title,
                    row_index,
                    header,
                )
            if _is_blank(value):
                _raise_config_error(
                    path,
                    "该行已填写其他字段，此字段不能留空",
                    worksheet.title,
                    row_index,
                    header,
                )

        rows.append((row_index, dict(zip(headers, values))))
    return rows


def _trimmed_header_values(worksheet) -> List[Any]:
    values = [
        worksheet.cell(row=1, column=column).value
        for column in range(1, worksheet.max_column + 1)
    ]
    while values and _is_blank(values[-1]):
        values.pop()
    return [value.strip() if isinstance(value, str) else value for value in values]


def _read_base_price_matrix(
    path: Path, worksheet
) -> Tuple[List[str], List[Tuple[int, Dict[str, Any]]]]:
    """读取「类别+材质」为行、尺寸为列的基础价格矩阵。"""

    _assert_no_data_merges(path, worksheet)
    headers = _trimmed_header_values(worksheet)
    if tuple(headers[:2]) != BASE_PRICE_FIXED_HEADERS or len(headers) < 3:
        _raise_config_error(
            path,
            "表头前两列必须是「类别名称 | 材质」，"
            "第三列起至少需要一个尺寸",
            worksheet.title,
            1,
        )

    sizes: List[str] = []
    seen_sizes: set[str] = set()
    for column, value in enumerate(headers[2:], start=3):
        cell = worksheet.cell(row=1, column=column)
        if cell.data_type == "f" or not isinstance(value, str):
            _raise_config_error(
                path,
                "尺寸表头必须是文本，不允许使用公式",
                worksheet.title,
                1,
                f"第{column}列",
            )
        size = _size(value, path, worksheet.title, 1, value)
        if size in seen_sizes:
            _raise_config_error(
                path,
                f"尺寸列「{size}」重复",
                worksheet.title,
                1,
                size,
            )
        seen_sizes.add(size)
        sizes.append(size)

    rows: List[Tuple[int, Dict[str, Any]]] = []
    column_count = 2 + len(sizes)
    for column in range(column_count + 1, worksheet.max_column + 1):
        for row in range(2, worksheet.max_row + 1):
            if not _is_blank(worksheet.cell(row=row, column=column).value):
                _raise_config_error(
                    path,
                    "价格单元格所在列缺少尺寸表头",
                    worksheet.title,
                    row,
                    f"第{column}列",
                )
    for row_index in range(2, worksheet.max_row + 1):
        cells = [
            worksheet.cell(row=row_index, column=column)
            for column in range(1, column_count + 1)
        ]
        values = [cell.value for cell in cells]
        if all(_is_blank(value) for value in values):
            continue
        for column, cell in enumerate(cells, start=1):
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                column_name = headers[column - 1]
                _raise_config_error(
                    path,
                    "数据单元格不允许使用公式",
                    worksheet.title,
                    row_index,
                    str(column_name),
                )
        for column_index, column_name in ((0, "类别名称"), (1, "材质")):
            if _is_blank(values[column_index]):
                _raise_config_error(
                    path,
                    "该行已填写价格，此字段不能留空",
                    worksheet.title,
                    row_index,
                    column_name,
                )
        if all(_is_blank(value) for value in values[2:]):
            _raise_config_error(
                path,
                "每行至少需要填写一个尺寸的单价",
                worksheet.title,
                row_index,
            )
        rows.append(
            (
                row_index,
                {
                    "类别名称": values[0],
                    "材质": values[1],
                    "价格": dict(zip(sizes, values[2:])),
                },
            )
        )
    return sizes, rows


def _read_pillow_price_matrix(path: Path, worksheet) -> List[Tuple[int, Dict[str, Any]]]:
    """读取尺寸和材质为行、六种枕芯类型为列的价格矩阵。"""

    _assert_no_data_merges(path, worksheet)
    _assert_headers(path, worksheet, PILLOW_MATRIX_HEADERS)
    rows: List[Tuple[int, Dict[str, Any]]] = []
    for row_index in range(2, worksheet.max_row + 1):
        cells = [
            worksheet.cell(row=row_index, column=column)
            for column in range(1, len(PILLOW_MATRIX_HEADERS) + 1)
        ]
        values = [cell.value for cell in cells]
        if all(_is_blank(value) for value in values):
            continue
        for column_name, cell in zip(PILLOW_MATRIX_HEADERS, cells):
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                _raise_config_error(
                    path,
                    "数据单元格不允许使用公式",
                    worksheet.title,
                    row_index,
                    column_name,
                )
        if _is_blank(values[0]):
            _raise_config_error(
                path,
                "尺寸不能留空",
                worksheet.title,
                row_index,
                "尺寸",
            )
        if _is_blank(values[1]):
            _raise_config_error(
                path,
                "枕芯材质不能留空",
                worksheet.title,
                row_index,
                "枕芯材质",
            )
        if all(_is_blank(value) for value in values[2:]):
            _raise_config_error(
                path,
                "每行至少需要填写一种枕芯类型的单价",
                worksheet.title,
                row_index,
            )
        rows.append((row_index, dict(zip(PILLOW_MATRIX_HEADERS, values))))
    return rows


def _casefold(value: str) -> str:
    return value.casefold()


def load_config_workbook(path: str | Path, *, check_lock: bool = True) -> CostConfig:
    """读取并严格校验成本配置工作簿。"""

    config_path = Path(path).resolve()
    if not config_path.is_file():
        raise ConfigWorkbookError(f"找不到成本配置文件：{config_path}")
    if check_lock:
        _ensure_not_in_use(config_path)

    try:
        workbook = load_workbook(config_path, read_only=False, data_only=False)
    except PermissionError as exc:
        raise ConfigWorkbookError(
            f"无法读取配置文件：{config_path}。"
            "请保存并关闭 Excel 后重试。"
        ) from exc
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ConfigWorkbookError(
            f"无法打开成本配置文件：{config_path}（{exc}）"
        ) from exc

    try:
        missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing_sheets:
            _raise_config_error(
                config_path,
                f"缺少工作表：{', '.join(missing_sheets)}",
            )

        global_rows = _read_rows(
            config_path,
            workbook[GLOBAL_SHEET],
            EXPECTED_HEADERS[GLOBAL_SHEET],
        )
        allowed_parameters = {"配置版本", "代发单价"}
        parameters: Dict[str, Tuple[int, Any]] = {}
        for row_index, row in global_rows:
            name = _clean_text(
                row["参数"], config_path, GLOBAL_SHEET, row_index, "参数"
            )
            if name not in allowed_parameters:
                _raise_config_error(
                    config_path,
                    f"不支持的参数「{name}」",
                    GLOBAL_SHEET,
                    row_index,
                    "参数",
                )
            if name in parameters:
                _raise_config_error(
                    config_path,
                    f"参数「{name}」重复",
                    GLOBAL_SHEET,
                    row_index,
                    "参数",
                )
            parameters[name] = (row_index, row["值"])

        for required_parameter in allowed_parameters:
            if required_parameter not in parameters:
                _raise_config_error(
                    config_path,
                    f"缺少参数「{required_parameter}」",
                    GLOBAL_SHEET,
                )

        version_row, version_value = parameters["配置版本"]
        version = _positive_integer(
            version_value,
            config_path,
            GLOBAL_SHEET,
            version_row,
            "值",
        )
        if version != CONFIG_VERSION:
            _raise_config_error(
                config_path,
                f"只支持配置版本 {CONFIG_VERSION}，当前为 {version}",
                GLOBAL_SHEET,
                version_row,
                "值",
            )

        dropship_row, dropship_value = parameters["代发单价"]
        dropship_unit_cost = float(
            _number(
                dropship_value,
                config_path,
                GLOBAL_SHEET,
                dropship_row,
                "值",
            )
        )

        category_rows = _read_rows(
            config_path,
            workbook[BASE_CATEGORY_SHEET],
            EXPECTED_HEADERS[BASE_CATEGORY_SHEET],
        )
        categories: Dict[str, Dict[str, Any]] = {}
        category_names_casefold: Dict[str, str] = {}
        priorities: Dict[int, str] = {}
        for row_index, row in category_rows:
            name = _clean_text(
                row["类别名称"],
                config_path,
                BASE_CATEGORY_SHEET,
                row_index,
                "类别名称",
            )
            priority = _positive_integer(
                row["优先级"],
                config_path,
                BASE_CATEGORY_SHEET,
                row_index,
                "优先级",
            )
            folded_name = _casefold(name)
            if folded_name in category_names_casefold:
                _raise_config_error(
                    config_path,
                    f"类别「{name}」重复",
                    BASE_CATEGORY_SHEET,
                    row_index,
                    "类别名称",
                )
            if priority in priorities:
                _raise_config_error(
                    config_path,
                    f"优先级 {priority} 已被类别「{priorities[priority]}」使用",
                    BASE_CATEGORY_SHEET,
                    row_index,
                    "优先级",
                )
            category_names_casefold[folded_name] = name
            priorities[priority] = name
            categories[name] = {
                "category": name,
                "keywords": [],
                "priority": priority,
                "products": [],
                "_products": {},
            }

        if not categories:
            _raise_config_error(config_path, "至少需要一个基础商品类别", BASE_CATEGORY_SHEET)

        keyword_rows = _read_rows(
            config_path,
            workbook[BASE_KEYWORD_SHEET],
            EXPECTED_HEADERS[BASE_KEYWORD_SHEET],
        )
        category_keyword_keys: Dict[str, set[str]] = {
            name: set() for name in categories
        }
        for row_index, row in keyword_rows:
            raw_category_name = _clean_text(
                row["类别名称"],
                config_path,
                BASE_KEYWORD_SHEET,
                row_index,
                "类别名称",
            )
            canonical_name = category_names_casefold.get(_casefold(raw_category_name))
            if canonical_name is None:
                _raise_config_error(
                    config_path,
                    f"类别「{raw_category_name}」未在「{BASE_CATEGORY_SHEET}」中定义",
                    BASE_KEYWORD_SHEET,
                    row_index,
                    "类别名称",
                )
            keyword = _clean_text(
                row["关键词"],
                config_path,
                BASE_KEYWORD_SHEET,
                row_index,
                "关键词",
            )
            key = _casefold(keyword)
            if key in category_keyword_keys[canonical_name]:
                _raise_config_error(
                    config_path,
                    f"类别「{canonical_name}」的关键词「{keyword}」重复",
                    BASE_KEYWORD_SHEET,
                    row_index,
                    "关键词",
                )
            category_keyword_keys[canonical_name].add(key)
            categories[canonical_name]["keywords"].append(keyword)

        _base_sizes, base_price_rows = _read_base_price_matrix(
            config_path, workbook[BASE_PRICE_SHEET]
        )
        base_price_keys: set[Tuple[str, str, str]] = set()
        for row_index, row in base_price_rows:
            raw_category_name = _clean_text(
                row["类别名称"],
                config_path,
                BASE_PRICE_SHEET,
                row_index,
                "类别名称",
            )
            canonical_name = category_names_casefold.get(_casefold(raw_category_name))
            if canonical_name is None:
                _raise_config_error(
                    config_path,
                    f"类别「{raw_category_name}」未在「{BASE_CATEGORY_SHEET}」中定义",
                    BASE_PRICE_SHEET,
                    row_index,
                    "类别名称",
                )
            material = _clean_text(
                row["材质"], config_path, BASE_PRICE_SHEET, row_index, "材质"
            )

            for size, raw_cost in row["价格"].items():
                if _is_blank(raw_cost):
                    continue
                cost = _number(
                    raw_cost, config_path, BASE_PRICE_SHEET, row_index, size
                )
                key = (_casefold(canonical_name), size, _casefold(material))
                if key in base_price_keys:
                    _raise_config_error(
                        config_path,
                        f"类别「{canonical_name}」、尺寸「{size}」、材质「{material}」重复",
                        BASE_PRICE_SHEET,
                        row_index,
                    )
                base_price_keys.add(key)

                category = categories[canonical_name]
                products_by_size = category["_products"]
                if size not in products_by_size:
                    product = {"尺寸": size, "价格": {}}
                    products_by_size[size] = product
                    category["products"].append(product)
                products_by_size[size]["价格"][material] = cost

        for category_name, category in categories.items():
            if not category["keywords"]:
                _raise_config_error(
                    config_path,
                    f"类别「{category_name}」至少需要一个关键词",
                    BASE_KEYWORD_SHEET,
                )
            if not category["products"]:
                _raise_config_error(
                    config_path,
                    f"类别「{category_name}」至少需要一条价格",
                    BASE_PRICE_SHEET,
                )
            category.pop("_products", None)

        pillow_rows = _read_pillow_price_matrix(
            config_path, workbook[PILLOW_PRICE_SHEET]
        )
        pillow_cost_data: Dict[str, Dict[str, int | float]] = {}
        pillow_keys: set[Tuple[str, str]] = set()
        pillow_row_keys: set[Tuple[str, str]] = set()
        for row_index, row in pillow_rows:
            size = _size(
                row["尺寸"], config_path, PILLOW_PRICE_SHEET, row_index, "尺寸"
            )
            material = _clean_text(
                row["枕芯材质"],
                config_path,
                PILLOW_PRICE_SHEET,
                row_index,
                "枕芯材质",
            )
            material_key = "" if _casefold(material) == _casefold("PP棉") else material
            row_key = (size, _casefold(material))
            if row_key in pillow_row_keys:
                _raise_config_error(
                    config_path,
                    f"尺寸「{size}」、枕芯材质「{material}」重复",
                    PILLOW_PRICE_SHEET,
                    row_index,
                )
            pillow_row_keys.add(row_key)

            for style_header, style_prefix in PILLOW_STYLE_PREFIXES.items():
                raw_cost = row[style_header]
                if _is_blank(raw_cost):
                    continue
                cost = _number(
                    raw_cost,
                    config_path,
                    PILLOW_PRICE_SHEET,
                    row_index,
                    style_header,
                )
                keyword = f"{style_prefix}{material_key}枕芯"
                key = (_casefold(keyword), size)
                if key in pillow_keys:
                    _raise_config_error(
                        config_path,
                        f"生成的关键词「{keyword}」、尺寸「{size}」重复",
                        PILLOW_PRICE_SHEET,
                        row_index,
                        style_header,
                    )
                pillow_keys.add(key)
                pillow_cost_data.setdefault(size, {})[keyword] = cost

        if not pillow_cost_data:
            _raise_config_error(config_path, "至少需要一条枕芯价格", PILLOW_PRICE_SHEET)

        moving_rows = _read_rows(
            config_path,
            workbook[MOVING_PRICE_SHEET],
            EXPECTED_HEADERS[MOVING_PRICE_SHEET],
        )
        moving_costs_data: List[Dict[str, Any]] = []
        moving_keys: set[str] = set()
        for row_index, row in moving_rows:
            keyword = _clean_text(
                row["关键词"],
                config_path,
                MOVING_PRICE_SHEET,
                row_index,
                "关键词",
            )
            folded_keyword = _casefold(keyword)
            if "yr" not in folded_keyword and "yt" not in folded_keyword:
                _raise_config_error(
                    config_path,
                    "义乳义臀关键词必须包含 yr 或 yt",
                    MOVING_PRICE_SHEET,
                    row_index,
                    "关键词",
                )
            if folded_keyword in moving_keys:
                _raise_config_error(
                    config_path,
                    f"关键词「{keyword}」重复",
                    MOVING_PRICE_SHEET,
                    row_index,
                    "关键词",
                )
            moving_keys.add(folded_keyword)
            moving_costs_data.append(
                {
                    "remark": keyword,
                    "cost": _number(
                        row["单价"],
                        config_path,
                        MOVING_PRICE_SHEET,
                        row_index,
                        "单价",
                    ),
                }
            )

        if not moving_costs_data:
            _raise_config_error(config_path, "至少需要一条义乳义臀价格", MOVING_PRICE_SHEET)

        other_rows = _read_rows(
            config_path,
            workbook[OTHER_PRICE_SHEET],
            EXPECTED_HEADERS[OTHER_PRICE_SHEET],
        )
        others_cost_data: List[Dict[str, Any]] = []
        other_keys: set[str] = set()
        for row_index, row in other_rows:
            keyword = _clean_text(
                row["关键词"],
                config_path,
                OTHER_PRICE_SHEET,
                row_index,
                "关键词",
            )
            folded_keyword = _casefold(keyword)
            if folded_keyword in other_keys:
                _raise_config_error(
                    config_path,
                    f"关键词「{keyword}」重复",
                    OTHER_PRICE_SHEET,
                    row_index,
                    "关键词",
                )
            other_keys.add(folded_keyword)
            others_cost_data.append(
                {
                    "remark": keyword,
                    "cost": _number(
                        row["单价"],
                        config_path,
                        OTHER_PRICE_SHEET,
                        row_index,
                        "单价",
                    ),
                }
            )

        if not others_cost_data:
            _raise_config_error(config_path, "至少需要一条其他成本", OTHER_PRICE_SHEET)

        price_data = sorted(
            categories.values(), key=lambda item: item["priority"]
        )
        return CostConfig(
            price_data=price_data,
            moving_costs_data=moving_costs_data,
            pillow_cost_data=pillow_cost_data,
            others_cost_data=others_cost_data,
            dropship_unit_cost=dropship_unit_cost,
            dropship_keywords=DROPSHIP_KEYWORDS,
            source_path=config_path,
        )
    finally:
        workbook.close()
