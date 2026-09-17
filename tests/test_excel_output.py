from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config_workbook import load_config_workbook
from cost_calculator import process_cost_detail_sheet, process_excel_file


ROOT = Path(__file__).resolve().parents[1]


class ExcelOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.config = load_config_workbook(
            ROOT / "成本配置模板.xlsx", check_lock=False
        )

    def test_reason_column_and_counts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            source = directory / "orders.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["订单编号", "卖家备注", "商家/店铺"])
            rows = [
                ["1", "B-1，桃皮绒50150枕套+枕芯50150", "店A"],
                ["2", "B-2，SAS50150枕套", "店A"],
                ["3", "ft枕芯", "店B"],
                ["4", "发海外", "店B"],
                ["5", "", "店C"],
                ["6", "三yr3250+起毛套50160", "店C"],
                ["7", "起毛套50160", "店C"],
            ]
            for row in rows:
                worksheet.append(row)
            workbook.save(source)
            workbook.close()

            config = self.config
            result = process_excel_file(
                str(source),
                config.price_data,
                config.moving_costs_data,
                config.pillow_cost_data,
                config.others_cost_data,
                dropship_unit_cost=config.dropship_unit_cost,
                dropship_keywords=config.dropship_keywords,
                output_dir=str(directory),
                overwrite=True,
            )

            self.assertTrue(result.success)
            self.assertEqual(7, result.total_count)
            self.assertEqual(2, result.matched_count)
            self.assertEqual(4, result.unmatched_count)
            self.assertEqual(1, result.overseas_count)

            output = load_workbook(result.output_path, data_only=False)
            detail = output["成本明细"]
            headers = [cell.value for cell in detail[1]]
            reason_index = headers.index("无法匹配原因说明") + 1
            other_index = headers.index("硅胶/电动成本") + 1
            self.assertEqual(other_index + 1, reason_index)
            self.assertEqual(50.0, detail.column_dimensions[detail.cell(1, reason_index).column_letter].width)

            reasons = [detail.cell(row, reason_index).value for row in range(2, 9)]
            self.assertIsNone(reasons[0])
            self.assertIn("不支持材质", reasons[1])
            self.assertIn("枕芯成本", reasons[2])
            self.assertEqual("发海外，成本需人工核对填写", reasons[3])
            self.assertEqual("卖家备注为空", reasons[4])
            self.assertIsNone(reasons[5])
            self.assertIn("未识别任何可计费项目", reasons[6])

            summary_row = len(rows) + 2
            self.assertIsNone(detail.cell(summary_row, reason_index).value)
            self.assertNotIn("无法匹配原因说明", [cell.value for cell in output["店铺统计"][1]])

            self.assertEqual(
                [item["reason"] for item in result.unmatched_details],
                [reasons[index] for index in (1, 2, 4, 6)],
            )
            output.close()

    def test_existing_reason_column_is_reused(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(
            [
                "商家/店铺",
                "成本",
                "代发成本",
                "总成本",
                "",
                "义乳/义臀成本",
                "枕芯成本",
                "硅胶/电动成本",
                "无法匹配原因说明",
            ]
        )
        worksheet.append(["店A"])

        result = process_cost_detail_sheet(
            worksheet,
            [0],
            [0],
            [0],
            [0],
            [0],
            ["新原因"],
            ["店A"],
            1,
        )

        self.assertIsNotNone(result)
        headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(1, headers.count("无法匹配原因说明"))
        reason_column = headers.index("无法匹配原因说明") + 1
        self.assertEqual("新原因", worksheet.cell(2, reason_column).value)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
