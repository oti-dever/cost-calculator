from __future__ import annotations

import copy
import io
import shutil
import tempfile
import threading
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from textual.widgets import DataTable, Input, Static, TextArea

import calculator_service
from calculator_service import CostCalculatorService
from config_workbook import ConfigWorkbookError, load_config_workbook
from cost_calculator import process_excel_file, recognize_note
from remark_tester import (
    RemarkTesterApp,
    RemarkTesterScreen,
    excel_cost_values,
    recognition_status,
)
from textual_app import CostCalculatorTextualApp


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "成本配置模板.xlsx"
VALID_NOTE = "B-1，桃皮绒50150枕套+枕芯50150"


def recognize(text: str, config=None, *, details=True):
    config = config or load_config_workbook(TEMPLATE, check_lock=False)
    return recognize_note(
        text, config.price_data, config.moving_costs_data, config.pillow_cost_data,
        config.others_cost_data, config.dropship_unit_cost, config.dropship_keywords,
        collect_details=details, pillow_metadata=config.pillow_metadata,
    )


class RecognitionDetailTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = load_config_workbook(TEMPLATE, check_lock=False)

    def test_collecting_details_preserves_results_and_reasons(self):
        notes = [
            VALID_NOTE, "B-1，B-2，桃皮绒50150枕套+枕芯50150+枕芯50150",
            "三yr3250+电动", "ft枕芯", "B-1，枕芯50150", "", "未知",
            "枕芯50150。未知", "枕芯50150。ft枕芯",
            "枕芯50150。三yr3250+起毛套50160", "发海外+" + VALID_NOTE,
        ]
        for note in notes:
            with self.subTest(note=note):
                plain = recognize(note, self.config, details=False)
                traced = recognize(note, self.config)
                self.assertEqual(plain.as_tuple(), traced.as_tuple())
                self.assertEqual(plain.failure_reasons, traced.failure_reasons)
                self.assertEqual([], plain.details)
                self.assertEqual(6, len(traced.as_tuple()))

    def test_all_cost_types_quantity_and_pillow_metadata(self):
        result = recognize(
            "B-1，B-2，桃皮绒50150枕套+枕芯50150+枕芯50150+三yr3250+电动",
            self.config,
        )
        self.assertEqual("识别成功", recognition_status(result))
        base = next(row for row in result.details if row.cost_type == "base_cost")
        self.assertEqual((2, 20, 40, "50*150", "桃皮绒"),
                         (base.quantity, base.unit_cost, base.subtotal, base.size, base.material))
        pillows = [row for row in result.details if row.cost_type == "pillow_cost"]
        self.assertEqual(2, len(pillows))
        self.assertTrue(all(row.material == "PP棉" and row.category == "标准等身" for row in pillows))
        for cost_type in ("base_cost", "dropship_cost", "moving_cost", "pillow_cost", "other_cost"):
            self.assertEqual(getattr(result, cost_type), sum(
                row.subtotal for row in result.details if row.cost_type == cost_type and row.included
            ))

    def test_normalization_segments_and_actual_matching_priority(self):
        result = recognize("  加重 FT 羽绒棉 枕芯 50 X 150 。\n三YR3250+电动+硅胶", self.config)
        self.assertEqual([], result.failure_reasons)
        pillow = next(row for row in result.details if row.cost_type == "pillow_cost")
        self.assertEqual((1, "加重ft羽绒棉枕芯", "羽绒棉", "加重分腿"),
                         (pillow.segment_index, pillow.keyword, pillow.material, pillow.category))
        others = [row for row in result.details if row.cost_type == "other_cost"]
        self.assertEqual(1, len(others))
        self.assertEqual(2, others[0].segment_index)
        self.assertIn("\n", others[0].segment_text)
        expected = next(item for item in self.config.others_cost_data if item["remark"] == others[0].keyword)
        self.assertEqual(expected["cost"], others[0].unit_cost)

    def test_partial_and_fatal_failures_mark_discarded_details(self):
        partial = recognize("枕芯50150。ft枕芯", self.config)
        self.assertEqual("部分匹配", recognition_status(partial))
        self.assertIsNone(partial.pillow_cost)
        self.assertTrue(all(not row.included for row in partial.details if row.cost_type == "pillow_cost"))
        self.assertTrue(all(row.included for row in partial.details if row.cost_type == "dropship_cost"))
        fatal = recognize("B-1，枕芯50150", self.config)
        self.assertEqual("识别失败", recognition_status(fatal))
        self.assertTrue(fatal.details)
        self.assertTrue(all(not row.included for row in fatal.details))
        self.assertEqual((None,) * 6, fatal.as_tuple())

    def test_zero_price_is_success_and_overseas_has_no_trace(self):
        config = copy.deepcopy(self.config)
        for category in config.price_data:
            for product in category["products"]:
                product["价格"] = {material: 0 for material in product["价格"]}
        zero = recognize("桃皮绒50150枕套", config)
        self.assertEqual("识别成功", recognition_status(zero))
        self.assertEqual(0, zero.details[0].subtotal)
        self.assertTrue(zero.details[0].included)
        overseas = recognize("发海外+枕芯50150", config)
        self.assertEqual("海外需人工核对", recognition_status(overseas))
        self.assertEqual((None,) * 6, overseas.as_tuple())
        self.assertEqual([], overseas.details)

    def test_export_matches_recognition_including_empty_totals(self):
        notes = [VALID_NOTE, "ft枕芯", "B-1，枕芯50150", "发海外", ""]
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "orders.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["卖家备注", "商家/店铺"])
            for note in notes:
                sheet.append([note, "测试店铺"])
            workbook.save(source)
            workbook.close()
            cfg = self.config
            with redirect_stdout(io.StringIO()):
                output = process_excel_file(
                    str(source), cfg.price_data, cfg.moving_costs_data, cfg.pillow_cost_data,
                    cfg.others_cost_data, cfg.dropship_unit_cost, cfg.dropship_keywords,
                )
            self.assertTrue(output.success)
            workbook = load_workbook(output.output_path)
            sheet = workbook["成本明细"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            columns = {"base_cost": "成本", "dropship_cost": "代发成本", "moving_cost": "义乳/义臀成本",
                       "pillow_cost": "枕芯成本", "other_cost": "硅胶/电动成本"}
            for row, note in enumerate(notes, 2):
                result = recognize(note, cfg)
                values = excel_cost_values(result)
                for attr, header in columns.items():
                    self.assertEqual(values[attr], sheet.cell(row, headers[header]).value, (note, attr))
                self.assertEqual("；".join(result.failure_reasons) or None,
                                 sheet.cell(row, headers["无法匹配原因说明"]).value)
                total = sheet.cell(row, headers["总成本"]).value
                if result.total_cost is None:
                    self.assertIsNone(total)
                else:
                    self.assertTrue(total.startswith("=IFERROR("))
                    self.assertEqual(result.total_cost, result.base_cost + result.dropship_cost)
            workbook.close()


class RemarkServiceTests(unittest.TestCase):
    def test_reloads_and_rejects_invalid_config_without_stale_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "成本配置.xlsx"
            shutil.copyfile(TEMPLATE, target)
            service = CostCalculatorService()
            with patch.object(calculator_service, "ensure_editable_config", return_value=target):
                self.assertEqual(26, service.recognize_note("枕芯50150").pillow_cost)
                workbook = load_workbook(target)
                sheet = workbook["枕芯价格"]
                for row in sheet.iter_rows(min_row=2):
                    if row[0].value == "150*50" and row[1].value == "PP棉":
                        row[2].value = 123
                workbook.save(target)
                workbook.close()
                self.assertEqual(123, service.recognize_note("枕芯50150").pillow_cost)
                lock = target.with_name("~$" + target.name)
                lock.touch()
                with self.assertRaises(ConfigWorkbookError):
                    service.recognize_note("枕芯50150")
                self.assertIsNone(service.config)
                lock.unlink()
                workbook = load_workbook(target)
                workbook["全局参数"].cell(2, 2, 999)
                workbook.save(target)
                workbook.close()
                with self.assertRaises(ConfigWorkbookError):
                    service.recognize_note("枕芯50150")
                self.assertIsNone(service.config)
                shutil.copyfile(TEMPLATE, target)
                self.assertEqual(26, service.recognize_note("枕芯50150").pillow_cost)


class RemarkInterfaceTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        config = patch("remark_tester.ensure_editable_config", return_value=TEMPLATE)
        config.start()
        self.addCleanup(config.stop)

    async def test_recognize_edit_clear_and_literal_text(self):
        result = recognize(VALID_NOTE)
        app = RemarkTesterApp()
        with patch.object(CostCalculatorService, "recognize_note", return_value=result) as service:
            async with app.run_test(size=(120, 48)) as pilot:
                screen = app.screen
                note = screen.query_one("#remark-input", TextArea)
                note.load_text(VALID_NOTE)
                await pilot.pause()
                await pilot.press("ctrl+enter")
                await app.workers.wait_for_complete()
                await pilot.pause()
                service.assert_called_once_with(VALID_NOTE)
                self.assertEqual("识别成功", str(screen.query_one("#remark-status", Static).renderable))
                self.assertEqual(len(result.details), screen.query_one(DataTable).row_count)
                note.load_text("[bold]未知[/bold]")
                await pilot.pause()
                self.assertEqual("输入已修改，请重新识别。", str(screen.query_one("#remark-status", Static).renderable))
                await pilot.click("#remark-clear")
                self.assertEqual("", note.text)
                self.assertIs(app.focused, note)
                self.assertEqual(0, screen.query_one(DataTable).row_count)
                await pilot.press("ctrl+enter")
                self.assertEqual("请输入备注。", str(screen.query_one("#remark-status", Static).renderable))
                self.assertEqual(1, service.call_count)

    async def test_error_recovery_button_and_open_config(self):
        app = RemarkTesterApp()
        with patch.object(CostCalculatorService, "recognize_note", side_effect=[ConfigWorkbookError("配置被占用"), recognize("ft枕芯")]), \
             patch("remark_tester.ensure_editable_config", return_value=TEMPLATE), \
             patch("remark_tester.open_with_default_app") as opener:
            async with app.run_test(size=(120, 48)) as pilot:
                screen = app.screen
                screen.query_one(TextArea).load_text("ft枕芯")
                await pilot.click("#remark-recognize")
                await app.workers.wait_for_complete()
                await pilot.pause()
                self.assertIn("配置被占用", str(screen.query_one("#remark-reasons", Static).renderable))
                self.assertFalse(screen.processing)
                self.assertFalse(screen.query_one("#remark-recognize").disabled)
                await pilot.click("#remark-recognize")
                await app.workers.wait_for_complete()
                await pilot.pause()
                self.assertEqual("部分匹配", str(screen.query_one("#remark-status", Static).renderable))
                await pilot.click("#remark-open-config")
                opener.assert_called_once_with(TEMPLATE)

    async def test_main_entry_preserves_inputs_results_and_service(self):
        app = CostCalculatorTextualApp()
        with patch.object(CostCalculatorService, "load_configs", return_value=None):
            async with app.run_test(size=(120, 48)) as pilot:
                app.query_one("#targets-input", Input).value = "保留的订单.xlsx"
                app.latest_results = ["existing result"]
                await pilot.click("#remark-test-btn")
                self.assertIsInstance(app.screen, RemarkTesterScreen)
                self.assertIsNot(app.service, app.screen.service)
                await pilot.press("escape")
                self.assertEqual("保留的订单.xlsx", app.query_one("#targets-input", Input).value)
                self.assertEqual(["existing result"], app.latest_results)

    async def test_main_entry_is_visible_at_standard_terminal_size(self):
        app = CostCalculatorTextualApp()
        with patch.object(CostCalculatorService, "load_configs", return_value=None):
            async with app.run_test(size=(80, 24)) as pilot:
                self.assertTrue(await pilot.click("#remark-test-btn"))
                self.assertIsInstance(app.screen, RemarkTesterScreen)

    async def test_standalone_close_exits(self):
        app = RemarkTesterApp()
        async with app.run_test(size=(100, 40)) as pilot:
            await pilot.click("#remark-close")
            self.assertTrue(app._exit)

    async def test_small_terminal_long_content_and_no_duplicate_submit(self):
        started = threading.Event()
        release = threading.Event()

        def slow_recognition(text):
            started.set()
            if not release.wait(5):
                raise TimeoutError("test worker timed out")
            return recognize(text)

        app = RemarkTesterApp()
        with patch.object(CostCalculatorService, "recognize_note", side_effect=slow_recognition) as service:
            async with app.run_test(size=(80, 24)) as pilot:
                screen = app.screen
                screen.query_one(TextArea).load_text(VALID_NOTE + "[literal]" * 50 + "。未知" * 20)
                await pilot.press("ctrl+enter")
                try:
                    self.assertTrue(started.wait(1))
                    self.assertTrue(screen.processing)
                    screen.action_recognize()
                    screen.action_close()
                    self.assertIs(screen, app.screen)
                    self.assertTrue(screen.query_one("#remark-recognize").disabled)
                finally:
                    release.set()
                await app.workers.wait_for_complete()
                await pilot.pause()
                self.assertEqual(1, service.call_count)
                self.assertEqual("\n".join(screen.last_result.failure_reasons),
                                 str(screen.query_one("#remark-reasons", Static).renderable))
                self.assertEqual(20, len(screen.last_result.failure_reasons))
                content = screen.query_one("#remark-content")
                self.assertGreater(content.max_scroll_y, 0)
                screen.query_one("#remark-details", DataTable).move_cursor(row=0)
                await pilot.pause()
                self.assertIn("[literal]", str(screen.query_one("#remark-selected", Static).renderable))


if __name__ == "__main__":
    unittest.main()
