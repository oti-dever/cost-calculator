from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import calculator_service
from calculator_service import CostCalculatorService


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "成本配置模板.xlsx"


class ServiceConfigReloadTests(unittest.TestCase):
    def test_each_batch_reloads_saved_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "成本配置.xlsx"
            shutil.copyfile(TEMPLATE, config_path)
            service = CostCalculatorService()

            with patch.object(
                calculator_service,
                "ensure_editable_config",
                return_value=config_path,
            ):
                first = service.load_configs()
                self.assertEqual(18, first.pillow_cost_data["120*40"]["枕芯"])

                workbook = load_workbook(config_path)
                worksheet = workbook["枕芯价格"]
                for row in range(2, worksheet.max_row + 1):
                    if (
                        worksheet.cell(row, 1).value == "120*40"
                        and worksheet.cell(row, 2).value == "PP棉"
                    ):
                        worksheet.cell(row, 3, 19)
                        break
                workbook.save(config_path)
                workbook.close()

                self.assertEqual([], service.process_targets([], None, False))
                self.assertEqual(
                    19, service.config.pillow_cost_data["120*40"]["枕芯"]
                )


if __name__ == "__main__":
    unittest.main()
